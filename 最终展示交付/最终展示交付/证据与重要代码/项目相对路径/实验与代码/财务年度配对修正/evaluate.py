"""Same-task before/after comparison and prespecified report-date subgroups."""
from pathlib import Path
import json,math,warnings
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.metrics import average_precision_score,roc_auc_score
from scipy.special import expit
import train
HERE=train.HERE;ROOT=train.ROOT;OUT=train.OUT

def dates(v):
    s=v.astype('string');good=s.str.match(r'^\d{4}-\d{2}-\d{2}(?:\s.*)?$',na=False)&~s.str.contains(r'-00(?:-|$)',na=False)
    return pd.to_datetime(s.where(good),errors='coerce',format='mixed')
def read_governance():
    path=next((ROOT/'大创数据集').rglob('CG_Ybasic.xlsx'))
    with warnings.catch_warnings():
        warnings.simplefilter('ignore',UserWarning);w=load_workbook(path,read_only=True,data_only=True);s=w.worksheets[0];s.reset_dimensions();it=s.iter_rows(values_only=True);head=next(it);next(it);next(it);g=pd.DataFrame(it,columns=head);w.close()
    g['Symbol']=pd.to_numeric(g.Stkcd,errors='coerce').astype('Int64').astype('string').str.zfill(6)
    g['economic_year']=pd.to_numeric(g.Reptdt.astype('string').str[:4],errors='coerce');g['publication']=dates(g.Annodt)
    g=g[g.Reptdt.astype('string').str.endswith('12-31',na=False)]
    period_end=dates(g.Reptdt);invalid=g.publication.notna()&period_end.notna()&(g.publication<period_end)
    g.loc[invalid,'publication']=pd.NaT
    return g.groupby(['Symbol','economic_year']).publication.min().reset_index(),path,int(invalid.sum())

def masks(year,s,gov,events):
    keys=s['keys'].reset_index(drop=True);n=len(keys)
    f=keys.merge(gov,on=['Symbol','economic_year'],how='left',validate='one_to_one')
    e=events[events.known_date<='2023-12-31'].groupby(['Symbol','occurrence_year']).known_date.min()
    f=f.merge(e.rename('first_known_date'),left_on=['Symbol','economic_year'],right_index=True,how='left',validate='one_to_one')
    cutoff=pd.Timestamp(f'{year+1}-06-30');available=f.publication.notna()&(f.publication<=cutoff)
    known=dates(f.first_known_date);already=known.notna()&(known<=cutoff)
    assert np.array_equal(known.notna().to_numpy(),s['y']==1)
    groups={'All':np.ones(n,bool),'ReportAvailable':available.to_numpy(),'Undisclosed':(available&~already).to_numpy()}
    f['cutoff']=str(cutoff.date());f['y_2023']=s['y'].astype(int);f['report_available']=available;f['annual_violation_already_disclosed']=already;f['eligible_undisclosed']=groups['Undisclosed']
    f.to_csv(OUT/f'eligibility_{year}.csv',index=False)
    counts=[]
    for name,mask in groups.items():
        counts.append({'year':year,'group':name,'rows':int(mask.sum()),'positives':int(s['y'][mask].sum()),'positive_fraction':float(s['y'][mask].mean()),'budget10':int(math.ceil(.1*mask.sum())),'supported_types':int((s['types'][mask].sum(0)>0).sum())})
    return groups,counts

def metrics(s,a,b,mask):
    ix=np.flatnonzero(mask);y=s['y'][ix];types=s['types'][ix];az=a[ix];bz=b[ix];n=len(y);k=int(math.ceil(.1*n));selected=np.argsort(-az,kind='stable')[:k]
    top=np.argsort(-bz,axis=1,kind='stable')[:,:2];pred=np.zeros_like(types);pred[np.arange(n)[:,None],top]=1
    hit=(pred[selected]*types[selected]).sum(0);den=types.sum(0);supported=den>0;pos=y==1
    ap=[average_precision_score(types[pos,j],bz[pos,j]) for j in range(10) if types[pos,j].sum()>0]
    return {'rows':n,'positives':int(y.sum()),'budget10':k,'found':int(y[selected].sum()),'recall':float(y[selected].sum()/y.sum()),'precision':float(y[selected].mean()),'overall_ap':float(average_precision_score(y,az)),'overall_auc':float(roc_auc_score(y,az)),
            'random_precision_reference':float(y.mean()),'precision_lift':float(y[selected].mean()/y.mean()),'type_macro_coverage':float((hit[supported]/den[supported]).mean()),'hint_precision':float(hit.sum()/(2*k)),
            'conditional_ap':float(np.mean(ap)),'conditional_logloss':float((np.logaddexp(0,bz[pos])-types[pos]*bz[pos]).mean()),'conditional_brier':float(((expit(bz[pos])-types[pos])**2).mean())}

def main():
    assert json.loads((OUT/'verification.json').read_text())['new_trainings']==24
    gov,path,invalid_publications=read_governance();events=pd.read_csv(ROOT/'标签成熟度与后续记录核查/results/all_event_years.csv.gz',dtype={'Symbol':'string'})
    sources=[path,ROOT/'标签成熟度与后续记录核查/results/all_event_years.csv.gz',HERE/'evaluate.py']
    before={str(p):train.prev.sha(p) for p in sources};rows=[];counts=[];prediction_sets={}
    original=pd.read_csv(OUT/'metrics.csv')
    for year in train.YEARS:
        p=train.pipeline.prepare(year,True);s=p['splits'][3];groups,c=masks(year,s,gov,events);counts+=c
        prediction_sets[year]={'symbols':s['keys'].Symbol.to_numpy().astype(str),'labels':s['y'],'groups':groups,'scores':{}}
        for model in train.MODELS:
            for seed in train.SEEDS:
                folder=OUT/str(year)/f'{model}_s{seed}'
                files={'Old':train.source(year,seed,model)/'predictions.npz','FrozenCorrected':folder/'frozen_corrected_predictions.npz','Corrected':folder/'predictions.npz'}
                for version,path in files.items():
                    raw=np.load(path);a=raw['a_logit'];b=raw['b_logit'];prediction_sets[year]['scores'][model,seed,version]=a
                    for group,mask in groups.items():
                        m=metrics(s,a,b,mask);rows.append({'year':year,'model':model,'seed':seed,'data_version':version,'group':group,**m})
                        if group=='All':
                            old=original[(original.year==year)&(original.model==model)&(original.seed==seed)&(original.data_version==version)].iloc[0]
                            for key in ['recall','precision','type_macro_coverage','overall_ap']:assert abs(m[key]-old[key])<1e-12
    df=pd.DataFrame(rows);df.to_csv(OUT/'group_metrics.csv',index=False);pd.DataFrame(counts).to_csv(OUT/'group_support.csv',index=False)
    numeric=[c for c in df.select_dtypes('number') if c not in ['year','seed']]
    annual=df.groupby(['year','model','data_version','group'])[numeric].mean().reset_index();annual.to_csv(OUT/'annual_summary.csv',index=False)
    summary=annual.groupby(['model','data_version','group'])[numeric].mean().reset_index();summary.to_csv(OUT/'summary.csv',index=False)
    base=df[df.data_version=='Old'];new=df[df.data_version=='Corrected'];merged=new.merge(base,on=['year','model','seed','group'],suffixes=('_new','_old'),validate='one_to_one')
    contrasts=merged[['year','model','seed','group']].copy()
    for key in ['recall','precision','found','overall_ap','type_macro_coverage','hint_precision']:contrasts[key+'_difference']=merged[key+'_new']-merged[key+'_old']
    contrasts.to_csv(OUT/'paired_changes.csv',index=False)
    # AP interval uses shared company weights across years and fixed model predictions.
    symbols=np.unique(np.concatenate([d['symbols'] for d in prediction_sets.values()]));rng=np.random.default_rng(20261014);samples=[]
    locations={y:np.searchsorted(symbols,d['symbols']) for y,d in prediction_sets.items()}
    for rep in range(400):
        weight=rng.multinomial(len(symbols),np.full(len(symbols),1/len(symbols)))
        for model in train.MODELS:
            for group in ['All','Undisclosed']:
                yearly=[]
                for year,d in prediction_sets.items():
                    mask=d['groups'][group];w=weight[locations[year]][mask];y=d['labels'][mask]
                    if (w*y).sum()==0:continue
                    delta=[]
                    for seed in train.SEEDS:
                        a=d['scores'][model,seed,'Corrected'][mask];b=d['scores'][model,seed,'Old'][mask]
                        delta.append(average_precision_score(y,a,sample_weight=w)-average_precision_score(y,b,sample_weight=w))
                    yearly.append(np.mean(delta))
                if len(yearly)==2:samples.append({'replicate':rep,'model':model,'group':group,'AP_difference':float(np.mean(yearly))})
    boot=pd.DataFrame(samples);interval=[]
    for (model,group),g in boot.groupby(['model','group']):
        actual=contrasts[(contrasts.model==model)&(contrasts.group==group)].overall_ap_difference.mean()
        interval.append({'model':model,'group':group,'AP_difference':float(actual),'lower':float(g.AP_difference.quantile(.025)),'upper':float(g.AP_difference.quantile(.975)),'replicates':len(g)})
    pd.DataFrame(interval).to_csv(OUT/'AP_intervals.csv',index=False)
    assert before=={str(p):train.prev.sha(p) for p in sources}
    train.prev.save_json(OUT/'evaluation_verification.json',{'fixed_cutoff':'June30 of economic year+1','report_publication_proxy':'CG_Ybasic annual report first publication on or after report period end','publication_before_period_end_rows_excluded':invalid_publications,'subgroup_selection_uses_dates_not_model_scores':True,'same_year_known_violation_excluded':True,'prior_year_violations_retained':True,'full_sample_metrics_reproduced':True,'source_hashes':before,'source_hashes_unchanged':True,'rows':len(df),'AP_cluster_bootstrap_replicates':400,'scope':'retrospective development evaluation; financial historical vintage not verified'})
    print(summary[['model','data_version','group','recall','precision','overall_ap']].to_string(index=False),flush=True)
if __name__=='__main__':main()
