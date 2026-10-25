"""只读原始Y，按公告快照过滤，按候选文书组分层抽取完整组。"""
from __future__ import annotations

import calendar
import json
import re
import tomllib
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from .storage import digest, file_digest, now, read_jsonl, write_json

PROJECT_DIR = Path(__file__).resolve().parent.parent
REQUIRED = {"ViolationID", "DeclareDate", "Activity", "ViolationYear",
            "ViolationTypeID", "IsViolated", "Symbol"}


def load_config(path: Path) -> dict:
    with path.open("rb") as stream:
        cfg = tomllib.load(stream)
    s, api = cfg["study"], cfg["api"]
    if not re.fullmatch(r"[A-Za-z0-9_.-]+", s["run_id"]) or s["run_id"] in {".", ".."}:
        raise ValueError("run_id只使用字母、数字、点、下划线和连字符")
    date.fromisoformat(s["snapshot_cutoff"])
    if type(s["sample_size"]) is not int or s["sample_size"] < 1:
        raise ValueError("sample_size必须为正整数")
    if not 0 <= s["holdout_fraction"] < 1:
        raise ValueError("holdout_fraction应在[0,1)内")
    if not api["base_url"].startswith("https://"):
        raise ValueError("API地址应使用https")
    for key in ("max_tokens", "max_attempts", "concurrency", "timeout_seconds"):
        if type(api[key]) is not int or api[key] < 1:
            raise ValueError(f"api.{key}必须为正整数")
    return cfg


def resolve(path: str) -> Path:
    return (PROJECT_DIR / path).resolve()


def run_path(cfg: dict) -> Path:
    return resolve(cfg["output"]["root"]) / cfg["study"]["run_id"]


def date_upper_bound(text: str) -> date | None:
    """00月/日按该期间末日过滤：仅纳入能确认在快照前的公告。"""
    match = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})(?:[ T]00:00:00)?", text.strip())
    if not match:
        return None
    year, month, day = map(int, match.groups())
    try:
        if month == 0:
            if day != 0:
                return None
            return date(year, 12, 31)
        if day == 0:
            day = calendar.monthrange(year, month)[1]
        return date(year, month, day)
    except (ValueError, calendar.IllegalMonthError):
        return None


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def strata_of(source: dict) -> list[str]:
    codes = set(re.findall(r"P\d{4}", source["ViolationTypeID"]))
    result = ["other_only" if codes == {"P2599"} else
              "other_with_types" if "P2599" in codes else "specific_types"]
    result.append("listed_company_y" if source["IsViolated"].strip().upper() == "Y" else
                  "listed_company_n_or_unknown")
    year = source["ViolationYear"]
    known = re.findall(r"(?:19|20)\d{2}", year)
    result.append("year_unknown" if not known else
                  "year_partial_unknown" if "N/A" in year.upper() else
                  "multiple_years" if len(set(known)) > 1 else "single_year")
    result.append("text_long" if len(source["Activity"]) > 1500 else "text_short")
    return result


def document_group(source: dict) -> str:
    num = source.get("DocumentNumber", "").strip()
    if num and num.upper() not in {"N/A", "NA", "NULL"}:
        # 编号有可能跨年重复，加入机构和公告日期；仅称候选文书组。
        key = [source.get("Promulgator", ""), num, source["DeclareDate"]]
    else:
        key = [source.get("Promulgator", ""), source.get("FileName", ""),
               source["DeclareDate"], source.get("Activity", "")]
    return "doc-" + digest(key)[:20]


def read_source(path: Path, sheet: str = "") -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        # 这份CSMAR文件的存储维度错误地标为一列。
        ws.reset_dimensions()
        rows = ws.iter_rows(values_only=True)
        header = [cell_text(v).strip() for v in next(rows)]
        if not REQUIRED.issubset(header):
            raise ValueError(f"原始Y缺少字段：{sorted(REQUIRED - set(header))}")
        out, seen = [], set()
        for row_num, row in enumerate(rows, start=2):
            if row_num <= 3:  # 第2、3行是中文标题和单位
                continue
            source = {key: cell_text(row[i] if i < len(row) else None)
                      for i, key in enumerate(header) if key}
            rid = source["ViolationID"].strip()
            if not rid:
                continue
            if rid in seen:
                raise ValueError(f"ViolationID重复：{rid}")
            seen.add(rid)
            out.append({"record_id": rid, "source_row": row_num,
                        "document_group": document_group(source),
                        "strata": strata_of(source), "source": source})
        return out
    finally:
        wb.close()


def sample_groups(records: list[dict], target: int, seed: int, holdout: float) -> list[dict]:
    groups = defaultdict(list)
    for rec in records:
        groups[rec["document_group"]].append(rec)
    ranked = sorted(groups, key=lambda g: digest([seed, g]))
    strata = sorted({s for r in records for s in r["strata"]})
    pools = {s: iter([g for g in ranked if any(s in r["strata"] for r in groups[g])])
             for s in strata}
    chosen, selected = set(), []
    while len(selected) < target:
        progress = False
        for s in strata:
            group = next((g for g in pools[s] if g not in chosen), None)
            if group is None:
                continue
            chosen.add(group)
            progress = True
            split = ("holdout" if int(digest([seed, "split", group])[:12], 16) / 16**12
                     < holdout else "development")
            selected.extend(dict(r, split=split) for r in groups[group])
            if len(selected) >= target:
                break
        if not progress:
            break
    return sorted(selected, key=lambda r: (r["document_group"], r["record_id"]))


def select_explicit_records(records: list[dict], selection: dict) -> list[dict]:
    """已审阅的扩批清单逐项选择；不重新随机抽样或补足其他记录。"""
    ids = selection.get('record_ids')
    if not isinstance(ids,list) or not ids or any(not isinstance(x,str) or not x for x in ids):
        raise ValueError('扩批清单record_ids须为非空字符串列表')
    if len(set(ids)) != len(ids):
        raise ValueError('扩批清单存在重复record_id')
    known={r['record_id']:r for r in records}
    missing=sorted(set(ids)-known.keys())
    if missing:
        raise ValueError(f'扩批清单含不在当前快照/有效文本范围的记录：{missing[:5]}')
    return [dict(known[rid],split='development') for rid in ids]


def prepare(cfg: dict) -> dict:
    from .extraction import EXTRACTION_SCHEMA, PROMPT_VERSION, SCHEMA_VERSION, SYSTEM_PROMPT

    dest = run_path(cfg)
    if (dest / "manifest.json").exists():
        return verify_run(cfg)
    records = read_source(resolve(cfg["source"]["path"]), cfg["source"].get("sheet", ""))
    cutoff = date.fromisoformat(cfg["study"]["snapshot_cutoff"])
    eligible, excluded = [], Counter()
    for record in records:
        source = record["source"]
        bound = date_upper_bound(source["DeclareDate"])
        reason = ("announcement_date_unknown_or_invalid" if bound is None else
                  "after_snapshot" if bound > cutoff else
                  "activity_empty" if not source["Activity"].strip() else None)
        if reason:
            excluded[reason] += 1
        else:
            eligible.append(record)
    s = cfg["study"]
    explicit=cfg.get('selection',{}).get('record_ids_file')
    if explicit:
        selection=json.loads(resolve(explicit).read_text(encoding='utf-8'))
        if selection.get('source_sha256') != file_digest(resolve(cfg['source']['path'])):
            raise ValueError('扩批清单与原Y版本不一致')
        selected=select_explicit_records(eligible,selection)
        if len(selected) != s['sample_size'] or s['holdout_fraction'] != 0:
            raise ValueError('显式扩批清单须与sample_size一致且holdout_fraction=0；沿用既有保留组')
    else:
        selected = sample_groups(eligible, s["sample_size"], s["seed"], s["holdout_fraction"])
    if not selected:
        raise ValueError("该快照下没有可用文本记录")
    dest.mkdir(parents=True, exist_ok=True)
    prepared = dest / "prepared.jsonl"
    prepared.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in selected),
                        encoding="utf-8")
    prompt_hash = digest([SYSTEM_PROMPT, EXTRACTION_SCHEMA])
    manifest = {"created_at": now(), "run_id": s["run_id"], "config": cfg,
                "config_hash": digest(cfg), "source_sha256": file_digest(resolve(cfg["source"]["path"])),
                "prepared_sha256": file_digest(prepared), "schema_version": SCHEMA_VERSION,
                "prompt_version": PROMPT_VERSION, "prompt_hash": prompt_hash,
                "source_rows": len(records), "eligible_rows": len(eligible),
                "excluded_counts": dict(excluded), "selected_rows": len(selected),
                "selected_document_groups": len({r["document_group"] for r in selected}),
                "split_counts": dict(Counter(r["split"] for r in selected)),
                "strata_counts": dict(Counter(t for r in selected for t in r["strata"])),
                "human_sampling": "skipped_per_user_prior_validation",
                "sampling": "seeded_strata_round_robin_whole_candidate_document_groups",
                "sampling_note": "样本是结构探索用分层样本；完整组使条数可略超目标，不作总体患病率估计。"}
    if explicit:
        manifest.update(sampling='reviewable_explicit_training_notice_batch',
            selection_file=explicit,selection_sha256=file_digest(resolve(explicit)),
            sampling_note='原Y匹配的训练期调用候选，具体事实仍须抽取后准入；既有pilot保留组未纳入。')
    audit = cfg["source"].get("prior_audit")
    if audit and resolve(audit).exists():
        manifest["prior_audit_reference"] = {"path": audit, "sha256": file_digest(resolve(audit))}
    write_json(dest / "schema.json", EXTRACTION_SCHEMA)
    (dest / "system_prompt.txt").write_text(SYSTEM_PROMPT, encoding="utf-8")
    write_json(dest / "manifest.json", manifest)
    return manifest


def verify_run(cfg: dict) -> dict:
    from .extraction import EXTRACTION_SCHEMA, SYSTEM_PROMPT

    dest = run_path(cfg)
    if not (dest / "manifest.json").exists():
        raise ValueError("请先运行 prepare")
    manifest = json.loads((dest / "manifest.json").read_text(encoding="utf-8"))
    if manifest["config_hash"] != digest(cfg):
        raise ValueError("配置已变化，请使用新的run_id准备独立版本")
    if manifest["prompt_hash"] != digest([SYSTEM_PROMPT, EXTRACTION_SCHEMA]):
        raise ValueError("提示词/schema已变化，请使用新的run_id")
    if manifest["prepared_sha256"] != file_digest(dest / "prepared.jsonl"):
        raise ValueError("prepared.jsonl发生变化，请恢复原样或使用新的run_id")
    return manifest
