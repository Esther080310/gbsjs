from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path("/Users/how/Desktop/分布式企业违规检测")
OUT = ROOT / "汇报材料" / "中期汇报20261023"
OUT.mkdir(parents=True, exist_ok=True)

SOURCE = Path("/Users/how/Desktop/中期汇报前.docx")
OUTPUT = OUT / "给老师的中期研究汇报_背景与扩展完整版.docx"
FLOW = OUT / "应用场景与模型闭环.png"

ASSET = ROOT / "汇报材料" / "中期汇报20261019" / "assets"
ATTR = ROOT / "实验与代码" / "主模型与隐私副模型归因展示" / "20261023-v1" / "figures"
UPDATE = ROOT / "实验与代码" / "年度更新机制收束" / "20261023-v1" / "figures"


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def apply_font(run, size=11, bold=False, color="000000", name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", *, style=None, bold_lead=None, align=None,
             size=11, color="222222", before=0, after=6, line=1.42,
             keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.first_line_indent = Cm(0.74) if style not in ("Title", "Heading 1", "Heading 2") else None
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        apply_font(r1, size=size, bold=True, color=color)
        r2 = p.add_run(text[len(bold_lead):])
        apply_font(r2, size=size, color=color)
    else:
        r = p.add_run(text)
        apply_font(r, size=size, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = level == 1 and text.startswith(("六、", "九、"))
    r = p.add_run(text)
    apply_font(r, size=16 if level == 1 else 13, bold=True, color="000000")
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Cm(0.74 + level * 0.45)
    p.paragraph_format.first_line_indent = Cm(-0.38)
    p.paragraph_format.space_after = Pt(3.5)
    p.paragraph_format.line_spacing = 1.35
    r = p.add_run(text)
    apply_font(r, size=10.8, color="222222")
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.4):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = 1
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, h in enumerate(headers):
        cell = hdr.cells[j]
        set_cell_shading(cell, "17324D")
        set_cell_margins(cell)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(h))
        apply_font(r, size=font_size, bold=True, color="FFFFFF")
        if widths:
            cell.width = Cm(widths[j])
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, v in enumerate(row):
            cell = cells[j]
            if i % 2 == 1:
                set_cell_shading(cell, "F4F7F9")
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.25
            r = p.add_run(str(v))
            apply_font(r, size=font_size, color="222222")
            if widths:
                cell.width = Cm(widths[j])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, path, caption, width=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(7)
    c.paragraph_format.keep_with_next = True
    r = c.add_run(caption)
    apply_font(r, size=9.3, color="555555")


def add_equation(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    apply_font(r, size=12, color="111111", name="Microsoft YaHei")


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r._r.append(fld_char1)
    r._r.append(instr)
    r._r.append(fld_char2)
    apply_font(r, size=9, color="666666")


def build():
    if not FLOW.exists():
        raise FileNotFoundError(f"缺少流程图：{FLOW}")
    doc = Document(SOURCE)
    clear_body(doc)

    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)
    add_page_number(section)

    for sty_name in ["Normal", "Body Text", "First Paragraph"]:
        if sty_name in doc.styles:
            sty = doc.styles[sty_name]
            sty.font.name = "Microsoft YaHei"
            sty._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            sty.font.size = Pt(11)
            sty.font.color.rgb = RGBColor(34, 34, 34)
    for n in [1, 2, 3]:
        sty_name = f"Heading {n}"
        if sty_name in doc.styles:
            sty = doc.styles[sty_name]
            sty.font.name = "Microsoft YaHei"
            sty._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            sty.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(5)
    r = title.add_run("基于多源分节点信息的上市公司违规风险排序")
    apply_font(r, size=22, bold=True, color="000000")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(12)
    r = sub.add_run("中期研究汇报修订稿  2026年10月")
    apply_font(r, size=11, color="555555")

    add_para(doc,
        "本项目面向监管资源有限、风险线索分散、原始数据存在使用边界的企业违规核查场景。我们将任务定义为年报和审计意见发布后的年度风险排序：三个节点保留市场、财务和治理等原始字段，只传递本地表示；融合端输出按核查预算截取的企业名单，并给出违规类型提示。当前已完成主模型、企业与违规类型评价、通信优化、经验隐私副模型、归因分析和年度更新机制。",
        size=11.2, color="222222", after=8)

    add_table(doc,
        ["阶段成果", "结果"],
        [
            ["风险排序", "2020—2022年前10%名单召回率为32.78%、30.39%和33.39%"],
            ["审计信息增量", "相对无审计模型，三年平均Recall@10%提高2.71个百分点"],
            ["类型提示", "名单命中企业的条件Top2命中率三年平均87.26%"],
            ["主模型通信", "六模型载荷3205降至899 B/企业，名单与Top2保持一致"],
            ["隐私副模型", "关键财务数值恢复成功率75.2%降至54.7%，召回损失约0.81个百分点"],
            ["年度维护", "实现增量回放、触发重训、保留旧版和版本回退的两期回放"],
        ], widths=[4.0, 12.0], font_size=9.5)

    add_heading(doc, "一、背景与现实需求", 1)
    add_heading(doc, "1. 监管首先要决定先查谁", 2)
    add_para(doc,
        "上市公司数量、公告文本和交易记录持续增长，而监管核查需要调取材料、访谈、核验合同与资金流，无法对所有企业采用同等强度。中国证监会随机抽查制度长期采用抽取部分企业开展检查的安排；监管实践也强调风险导向选择检查对象。对模型而言，现实任务不是代替监管机关认定违规，而是在固定核查预算下把更可能存在问题的企业排到名单前部。")
    add_para(doc,
        "国家层面的资本市场政策进一步强调提高监管科技水平和跨部门信息共享。国务院《关于加强监管防范风险推动资本市场高质量发展的若干意见》提出加强监管协同和数据信息共享；关于资本市场财务造假综合惩防的文件指出，财务造假具有系统性、隐蔽性和复杂性，需要综合使用信用信息、资金流水和科技手段发现线索。这为“多源信息形成风险线索”的应用方向提供了政策依据。")

    add_heading(doc, "2. 违规线索分散在不同主体和业务环节", 2)
    add_para(doc,
        "公开案例显示，违规往往持续多年，并可能同时涉及财务报表、合同、客户供应商、银行流水、关联关系和审计工作。华微电子以预付设备款掩盖资金占用、第三方协助上市公司形成虚假资金闭环等案例说明，单份年报中的数字可能形式正常，只有与交易关系和资金流向交叉核验后，风险才会显现。市场、财务、治理、审计以及机构内部记录因而具有互补价值。")

    add_heading(doc, "3. 信息互补不等于原始数据可以无限集中", 2)
    add_para(doc,
        "本项目使用CSMAR公开数据完成方法实验，但现实协作中可能接入银行流水、审计底稿、内部评级、客户供应商关系或未公开核查线索。这些数据分别受到权限、用途、商业秘密和数据安全要求约束。《关于构建数据基础制度更好发挥数据要素作用的意见》提出在保护个人信息和公共安全的前提下探索“原始数据不出域、数据可用不可见”。金融领域的联邦学习和多方安全计算研究也说明，多源协作与信息保护需要同时设计。")
    add_para(doc,
        "分节点建模只解决原始字段不汇集的问题，并不自动等于安全。中间表示、梯度和最终输出仍可能泄露信息。因此本项目把“预测是否有效”“传输多少信息”“表示能否被攻击者还原”分成三条证据链，分别检验。")

    add_heading(doc, "4. 通报滞后要求模型能够持续更新", 2)
    add_para(doc,
        "违规发生、年报披露、审计意见和监管通报通常不在同一时点。既有研究对证监会2021年行政处罚的统计显示，信息披露违法从行为发生到处罚作出平均约44.9个月；本项目数据采用另一种口径，从违规发生年度结束到数据库记录的中位时间约15.8个月。两组数字口径不同，但都说明评分时尚未出现通报不能被理解为企业一定合规。模型输出因此是优先核查顺序，并需要在新通报到达后修正历史标签。")

    add_figure(doc, FLOW, "图1 现实核查流程与项目各模块的对应关系", width=6.55)

    add_heading(doc, "二、研究任务与数据边界", 1)
    add_para(doc,
        "样本单位为企业年度。评分时点设在年报和审计意见发布后，使用当时能够取得的市场、财务、治理、审计和企业行业信息，预测后续被记录的违规风险。融合端按风险分数排序，并按核查预算取前k%的企业。Recall@10%表示这张前10%名单覆盖全部记录正例的比例；Precision@10%表示名单中记录正例的比例。")
    add_para(doc,
        "企业行业作为条件输入和分层评价变量，用于观察不同经营背景下的模型表现；违规类型作为多标签输出，为名单内企业提供Top2核查线索。类型提示不代替调查结论，同一企业也可能对应多种违规。")
    add_para(doc,
        "当前实验以公开数据构造三个物理节点：X1为市场与交易信息，X2为财务报表和经济指标，X3为治理、事件及审计来源。公开数据使我们能够复现实验和设置集中式对照；其作用是验证分节点算法、通信接口和攻击评价方法。真实机构部署仍需明确各方掌握的字段、处理权限、标签归属和攻击者能力。")

    add_heading(doc, "三、主模型与方法创新", 1)
    add_para(doc,
        "市场、财务和治理节点分别形成16维、88维和21维表示。审计类别被编码为8维条件信息。模型不把审计意见简单当作第四组普通特征，而是用它调整财务表示的尺度和偏移；没有匹配审计记录时，调制支路关闭。")
    add_equation(doc, "h̃₂ = h₂ + r[γ(c) ⊙ h₂ + β(c)]")
    add_para(doc,
        "其中h₂是财务节点表示，c是审计条件编码，r是审计记录匹配标志。γ(c)和β(c)分别改变财务潜在表示的尺度和偏移。这个结构表达了一个明确假设：同一组财务状态在不同审计意见下可能具有不同风险含义。它调节的是88维潜在表示，不能直接解释为给某一项原始财务指标赋权。")
    add_figure(doc, ASSET / "02_FiLM-L模型结构_专业版.png", "图2 FiLM-L三节点编码与审计条件调节结构", width=6.4)
    add_para(doc,
        "我们设置U0、UA7、FiLM-L和集中式网络四个对照。U0不使用审计信息；UA7把同样的审计类别作为普通输入；FiLM-L使用条件调节；集中式网络用于观察分节点表示保留了多少排序能力。这样可以把“新增审计信息是否有用”和“条件调节结构是否优于普通接入”分开回答。")

    add_heading(doc, "四、风险排序结果", 1)
    add_para(doc,
        "训练、轮次选择和评价均按时间顺序进行，尺度参数只在训练数据上拟合。最终模型采用六个独立训练模型的校准logit平均。各模型使用相同随机种子和集成规则；模型差值通过企业聚类配对bootstrap评价，同一企业跨年度共用抽样。")
    add_table(doc,
        ["年度", "样本数", "记录正例", "前10%名单", "命中数", "Recall@10%", "Precision@10%", "AP"],
        [
            ["2020", "4292", "787", "430", "258", "32.78%", "60.00%", "0.4927"],
            ["2021", "4794", "747", "480", "227", "30.39%", "47.29%", "0.4108"],
            ["2022", "5199", "629", "520", "210", "33.39%", "40.38%", "0.3571"],
        ], widths=[1.6,1.8,1.9,2.0,1.6,2.2,2.4,1.7], font_size=8.8)
    add_figure(doc, ASSET / "03_模型效果对照.png", "图3 主模型与无审计、普通审计接入及集中式网络的比较", width=6.35)
    add_para(doc,
        "UA7相对U0的三年平均Recall@10%提高2.71个百分点，95%区间为[1.68, 3.67]，AP平均提高4.38个百分点，说明审计类别提供了原有三节点之外的增量。FiLM-L相对UA7的Recall@10%平均低0.15个百分点，区间为[−0.87, 0.64]，现有真实数据没有显示条件调节比普通接入具有稳定精度优势。我们保留FiLM-L，是因为它把审计与财务的关系写进模型结构，并为后续归因和保护提供了清晰接口。")
    add_para(doc,
        "2020—2021年，FiLM-L相对集中式网络的Recall@10%平均高0.17个百分点，区间为[−1.21, 1.21]。这里可以说平均点值接近、未检出稳定差异，不能表述为已经证明统计等价。")
    add_figure(doc, ASSET / "05_前10前5前1名单密度.png", "图4 2020年前10%、前5%和前1%名单中的记录正例密度", width=6.25)
    add_para(doc,
        "以2020年为例，全体企业记录正例比例为18.3%；前10%名单的Precision为60.0%，前5%为72.6%，前1%为93.0%。名单越短，正例越集中，但覆盖总数会下降。前10%因此作为主要比较口径，前5%和前1%用于展示不同核查预算下的选择。")

    add_heading(doc, "五、企业类别与违规类型", 1)
    add_para(doc,
        "分类拓展包含两个不同问题。企业类别采用行业编码作为输入条件和分层变量，帮助融合端区分不同经营背景；违规类型采用十类多标签输出，为已进入名单的企业提供Top2提示。企业行业不是本项目新增的预测标签，违规类型才是模型输出。")
    add_table(doc,
        ["年度", "名单命中企业的条件Top2", "常见两类固定提示", "端到端类型宏覆盖"],
        [
            ["2020", "87.60%", "84.88%", "7.32%"],
            ["2021", "88.55%", "87.22%", "6.43%"],
            ["2022", "85.65%", "86.60%", "6.59%"],
            ["三年平均", "87.26%", "86.24%", "6.78%"],
        ], widths=[2.7,4.4,4.3,4.2], font_size=9.3)
    add_para(doc,
        "条件Top2回答“已命中的违规企业，其真实类型是否出现在前两项提示中”；端到端宏覆盖还把企业是否进入前10%名单计算在内，并对十种类型等权。后者较低，说明长尾类型仍是当前系统的主要限制。2022年数据包行业字段全部为未知值，因此行业比较只使用2020—2021年有编码样本，不将未知值当作企业类别。")
    add_figure(doc, ATTR / "02_行业与违规类型.png", "图5 不同行业和违规类型下的财务指标替换敏感性", width=6.4)

    add_heading(doc, "六、通信量优化", 1)
    add_para(doc,
        "分节点计算的通信重点不是简单追求向量更短，而是在不改变最终核查名单和类型提示的条件下，减少每家企业需要传输的数值。主模型采用六个独立编码器，因此正式账本按六套节点消息计算。我们对市场、财务、治理和审计控制码分别选择数值精度，先在无标签候选中生成81种组合，再用冻结评价筛选。")
    add_table(doc,
        ["方案", "六模型载荷", "相对FP32", "三年名单", "名单内Top2"],
        [
            ["原始FP32接口", "3205 B/企业", "基准", "原结果", "原结果"],
            ["市场FP16，财务/治理/审计INT8", "899 B/企业", "下降71.95%", "逐项相同", "逐项相同"],
        ], widths=[5.1,3.0,2.8,2.5,2.5], font_size=9.2)
    add_para(doc,
        "三个物理节点可以单轮并行上行，X3节点承载治理表示和审计控制码两个逻辑块。899 B是同一FiLM-L预测器在不同消息精度下的载荷，不包含TCP/TLS等协议开销，也不等于完整训练生命周期的网络流量。集中式模型需要先把原始字段集中，因而不与这一消息载荷作同口径比较。")

    add_heading(doc, "七、归因分析与隐私保护副模型", 1)
    add_heading(doc, "1. 为什么分节点后仍需检验隐私", 2)
    add_para(doc,
        "融合端虽然看不到原始财务表，但可能根据88维财务表示推断具体数值或企业属性。我们把融合端设为遵守协议但会保存和分析消息的攻击者，使用相同辅助数据和攻击预算，分别检验13项关键财务数值的容差内恢复成功率与属性推断AUC。")
    add_heading(doc, "2. 隐私—效用取舍", 2)
    add_para(doc,
        "在与主任务兼容的共享编码器框架下，我们比较无防护S0、距离相关约束dCor和随机信息瓶颈MID。MID通过随机潜变量限制表示携带的信息，作为可选副模型，不替换正式FiLM-L。五组结果是五次独立稳定性复检，不是部署时同时发送五份消息。")
    add_table(doc,
        ["模型", "Recall@10%", "AP", "财务数值猜中率", "属性攻击AUC"],
        [
            ["S0 无防护", "30.13%", "0.5768", "75.2%", "0.960"],
            ["dCor", "29.34%", "0.5628", "65.3%", "0.936"],
            ["MID 0.1", "29.32%", "0.5714", "54.7%", "0.909"],
        ], widths=[3.1,2.8,2.5,3.7,3.2], font_size=9.3)
    add_para(doc,
        "相对无防护模型，MID使关键财务数值恢复成功率下降20.5个百分点，平均Recall@10%下降约0.81个百分点。属性AUC仍为0.909，说明风险没有被消除。这里的结论是特定攻击协议下的经验抑制与代价，不是差分隐私保证，也不代表真实机构已经达到安全验收标准。")
    add_figure(doc, ATTR / "04_字段依赖与隐私保护.png", "图6 任务依赖与财务字段恢复难度的并列评价", width=6.35)
    add_heading(doc, "3. 归因如何指导保护", 2)
    add_para(doc,
        "四因素Shapley分解显示，主模型对X1市场、X2财务、X3治理事件和审计调制分支的平均绝对贡献占比分别为20.17%、41.93%、2.85%和35.06%。这说明隐私检测应重点检查财务表示和审计条件，而不是平均保护所有节点。MID仍主要使用财务与审计信息，但X3治理事件的相对贡献下降。归因结果用于解释模型依赖和选择攻击对象，不用于证明财务指标对违规的因果作用。")
    add_para(doc,
        "MID采用FP16传输时，单模型载荷由545降至273 B/企业，五组前10%名单和名单内Top2均保持，最大类型分数误差约0.000254。由于攻击器只在FP32消息上重新训练，报告中分别陈述通信兼容和隐私攻击结果，不把两者合并成已经验证的FP16安全保证。")

    add_heading(doc, "八、年度更新与版本回退", 1)
    add_para(doc,
        "新通报会把旧企业年度从未记录正例修正为已记录正例，也可能补充新的违规类型。若每年都从头训练三个节点，计算和通信成本较高；若只微调，又可能在政策或业务规律变化后保留过时表示。项目因此实现条件更新机制：平时冻结节点编码器，复用企业年度表示，只更新融合端，并把新增或修正样本与历史样本混合回放；保留集退化时评估全量重训；全量候选仍不通过时保留旧版本。")
    add_equation(doc, "min L更新(θ; φ旧) + λL历史回放(θ; φ旧)")
    add_para(doc,
        "标签变化不要求重新编码同一企业年度；新增年度只计算新键。缓存键同时包含编码器、预处理和数据版本，任一变化都会使缓存失效。模型版本记录父版本、标签截止日、文件哈希和编码器哈希，因此可以核对后回退。重大政策变化不由模型自动猜测，而是登记生效日期，在新的政策后监督达到预设数量时优先进行全量更新评估。")
    add_table(doc,
        ["后续评价年", "不更新", "每期全量", "连续增量回放", "监测后选择"],
        [
            ["2019 Recall@10%", "29.11%", "30.67%", "30.05%", "30.42%"],
            ["2020 Recall@10%", "30.20%", "32.61%", "31.05%", "31.34%"],
            ["两期累计计算时间", "0", "16.65秒", "0.37秒", "8.70秒"],
        ], widths=[4.0,2.7,2.7,3.3,3.2], font_size=9.2)
    add_figure(doc, UPDATE / "01_两期更新策略.png", "图7 两期连续回放中的效果与计算成本", width=6.35)
    add_para(doc,
        "连续增量回放在第一期与全量重训相差0.62个百分点，第二期相差1.57个百分点，说明微调不能长期无条件替代重训。监测策略实际走通过增量、全量和保留三种分支，两期累计时间约为每期全量重训的52.25%。这项结果支持“有监测和回退的条件更新”，不支持“每年微调一定足够”。")

    add_heading(doc, "九、研究探索与最终选型", 1)
    add_para(doc,
        "项目早期围绕经济结构压缩、表格学习器、不均衡学习、近邻案例和延迟标签修正开展了多组对照。多数方法只在个别年度或随机种子改善，无法跨年度复现，因此没有叠加到正式模型。保留这些结果的作用是解释选型边界，而不是把无效模块包装成最终方法。")
    add_table(doc,
        ["研究假设", "主要检验", "对最终模型的影响"],
        [
            ["经济关系有助于低维压缩", "财务分组、比率、交互选择与模拟", "真实数据无稳定增量；模拟中只在强交互与正确先验下受益"],
            ["复杂表格学习器提高X2能力", "TabNet、VIME、NeuMiss等迁移", "未跨年度稳定复现，未替换X2编码器"],
            ["新不均衡方法改善名单前部", "代价敏感、重采样及迁移方法", "没有稳定超过冻结参照"],
            ["历史案例提供模型外增量", "近邻风险与联合案例学习", "未发现稳定增量"],
            ["延迟比例可由历史队列修正", "重要性加权、趋势外推、自训练", "记录速度存在时期漂移，直接加权会破坏排序"],
            ["审计信息补充原三节点", "U0、匹配标志、UA7与FiLM-L", "三年均改善，进入正式模型"],
        ], widths=[4.2,5.1,6.0], font_size=8.8)
    add_para(doc,
        "经济结构模拟按真实数据维度、相关性和类别比例构造40组数据，共完成360次训练。当模拟数据存在与经济分组一致的强交互时，24维经济结构压缩相对同维随机压缩的Recall@10%提高1.68个百分点，95%区间为[0.98, 2.38]；先验错位或交互较弱时优势减弱或消失。这项实验说明方法的适用条件，但不替代真实数据结论。")

    add_heading(doc, "十、阶段成果与汇报结论", 1)
    add_para(doc,
        "项目目前形成了一套完整的企业违规优先核查原型：以监管预算决定名单长度，以分节点编码处理多源信息，以审计条件调节财务表示，以多标签输出提供违规类型线索；在此基础上，进一步验证了名单保持下的通信压缩、归因引导的经验隐私保护和通报到达后的年度更新。")
    add_para(doc, "当前创新可以分为两个层次：", bold_lead="当前创新可以分为两个层次：")
    add_bullet(doc, "核心方法：将审计判断作为财务表示的条件信息，在原始字段不汇集的三节点结构中完成风险排序和类型提示，并用同信息对照拆分信息增量与结构作用。")
    add_bullet(doc, "系统拓展：围绕同一核查任务，建立决策保持的通信精度选择、以模型依赖为线索的攻击评价与MID副模型、以及有监测和回退的年度维护机制。")
    add_para(doc,
        "主模型的实际价值体现在排序前端的风险富集，而不是给所有企业作确定裁决。2020年前10%名单将记录正例比例从全体18.3%提高到60.0%；三个测试年度前10%名单分别覆盖32.78%、30.39%和33.39%的记录正例。对没有真实合作机构数据的当前阶段，这些结果支持方法与接口原型；若进入机构合作，还需根据实际数据权限、攻击者能力和核查成本重新确定验收标准。")
    add_para(doc,
        "本阶段停止继续搜索新主干。后续工作集中在中期汇报、材料归档和复现整理。正式主模型保持FiLM-L，MID作为可选隐私副模型，年度更新作为维护拓展；三者的样本和评价口径分别陈述。")

    add_heading(doc, "主要依据", 1)
    refs = [
        "[1] 国务院. 关于加强监管防范风险推动资本市场高质量发展的若干意见（国发〔2024〕10号）.",
        "[2] 国务院办公厅等. 关于进一步做好资本市场财务造假综合惩防工作的意见, 2024.",
        "[3] 中共中央、国务院. 关于构建数据基础制度更好发挥数据要素作用的意见, 2022.",
        "[4] 中国证监会. 中国证监会随机抽查事项清单（证监会公告〔2021〕1号）.",
        "[5] 黄江东, 施蕾. 证券监管执法的若干典型问题及完善建议——基于证监会2021年345例行政处罚的实证研究. 投资者, 2022.",
        "[6] 中国证监会及派出机构公开的上市公司财务造假、资金占用与信息披露违法案例材料.",
        "[7] 中国人民银行. 多方安全计算金融应用技术规范：JR/T 0196—2020.",
        "[8] Yang Q, Liu Y, Chen T, et al. Federated Machine Learning: Concept and Applications. ACM TIST, 2019.",
        "[9] Khan A, ten Thij M, Wilbik A. Vertical federated learning: a structured literature review. Knowledge and Information Systems, 2025.",
        "[10] Zhu L, Liu Z, Han S. Deep Leakage from Gradients. NeurIPS, 2019.",
        "[11] Fu C, Zhang X, Ji S, et al. Label Inference Attacks Against Vertical Federated Learning. USENIX Security, 2022.",
        "[12] Abowd J M, Schmutte I M. An Economic Analysis of Privacy Protection and Statistical Accuracy as Social Choices. AER, 2019.",
        "[13] 王小燕, 冮建伟, 方匡南. 隐私保护下基于多源域迁移学习的信用风险预警研究. 统计研究, 2026.",
        "[14] BIS Innovation Hub. Project Aurora: The Power of Data, Technology and Collaboration to Combat Money Laundering Across Institutions and Borders, 2023.",
    ]
    for ref in refs:
        add_para(doc, ref, size=9.3, color="333333", after=3.2, line=1.25)

    # Remove first-line indent from captions, tables, title-like paragraphs where appropriate.
    for p in doc.paragraphs:
        if p.alignment == WD_ALIGN_PARAGRAPH.CENTER or p.style.name.startswith("Heading") or p.style.name == "Title":
            p.paragraph_format.first_line_indent = None

    doc.core_properties.title = "基于多源分节点信息的上市公司违规风险排序 中期研究汇报"
    doc.core_properties.subject = "背景、主模型、分类、通信、隐私、归因与年度更新"
    doc.core_properties.author = "项目组"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
